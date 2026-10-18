
import json
import re
import sys
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from sentence_transformers import SentenceTransformer, util


# ---------------------------------------------------------
# 1. Skill catalogue
# Add more skills and aliases as you test different resumes.
# ---------------------------------------------------------

SKILLS = {
    "Manual Testing": [
        "manual testing", "manual tester", "functional testing",
        "regression testing", "system testing"
    ],
    "Test Case Design": [
        "test case design", "test cases", "test scenarios",
        "test case preparation"
    ],
    "Test Planning": [
        "test planning", "testing strategy", "test plan"
    ],
    "Defect Management": [
        "defect management", "bug tracking", "defect tracking",
        "bug reporting", "defect reporting"
    ],
    "Automation Testing": [
        "automation testing", "test automation",
        "automated testing", "automated test scripts"
    ],
    "Selenium": ["selenium", "selenium webdriver"],
    "API Testing": [
        "api testing", "rest api testing", "postman",
        "restful api testing"
    ],
    "SQL": [
        "sql", "mysql", "postgresql", "sql queries",
        "database queries"
    ],
    "Python": ["python", "python programming"],
    "Java": ["java programming", "java"],
    "JavaScript": ["javascript", "js programming"],
    "Git": ["git", "github", "version control"],
    "CI/CD": [
        "ci/cd", "continuous integration",
        "continuous delivery", "continuous deployment"
    ],
    "Agile": ["agile", "agile methodology", "scrum"],
    "Jira": ["jira", "atlassian jira"],
    "Business Analysis": [
        "business analysis", "business analyst",
        "requirements gathering", "requirements analysis"
    ],
    "Requirements Documentation": [
        "requirements documentation", "business requirements",
        "functional requirements", "brd", "frd"
    ],
    "UAT": [
        "user acceptance testing", "uat",
        "user acceptance test"
    ],
    "Banking Domain Knowledge": [
        "banking domain", "banking operations",
        "core banking", "retail banking",
        "banking applications", "financial services"
    ],
    "Risk and Compliance": [
        "risk management", "regulatory compliance",
        "compliance testing", "audit compliance"
    ],
    "Data Analysis": [
        "data analysis", "data analytics",
        "data visualization", "data interpretation"
    ],
    "Power BI": ["power bi", "powerbi"],
    "Microsoft Excel": [
        "microsoft excel", "ms excel", "advanced excel"
    ],
    "Project Management": [
        "project management", "project planning",
        "project coordination"
    ],
    "Stakeholder Management": [
        "stakeholder management", "stakeholder communication",
        "stakeholder engagement"
    ],
    "Team Leadership": [
        "team leadership", "led a team",
        "managed a team", "team management"
    ],
    "Cloud Computing": [
        "cloud computing", "cloud infrastructure",
        "amazon web services", "aws", "microsoft azure"
    ],
    "Machine Learning": [
        "machine learning", "ml models",
        "model training", "predictive modeling"
    ],
    "Natural Language Processing": [
        "natural language processing", "nlp",
        "text classification", "named entity recognition"
    ],
}


# ---------------------------------------------------------
# 2. Read resume formats
# ---------------------------------------------------------

def read_resume(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        with fitz.open(path) as pdf:
            return "\n".join(page.get_text() for page in pdf)

    if suffix == ".docx":
        document = Document(path)
        paragraphs = [p.text for p in document.paragraphs]

        # Also read tables, which commonly contain resume details.
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(paragraphs)

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    raise ValueError(f"Unsupported file type: {suffix}")


# ---------------------------------------------------------
# 3. Exact/alias skill extraction with evidence
# ---------------------------------------------------------

def normalize(text: str) -> str:
    text = text.lower()
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9+#./-]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def extract_explicit_skills(text: str) -> list:
    normalized_text = normalize(text)
    found = []

    for canonical_name, aliases in SKILLS.items():
        matched_alias = None

        for alias in aliases:
            normalized_alias = normalize(alias)

            # Avoid matching a short skill inside a larger word.
            pattern = (
                r"(?<![a-z0-9])"
                + re.escape(normalized_alias)
                + r"(?![a-z0-9])"
            )

            match = re.search(pattern, normalized_text)

            if match:
                matched_alias = alias
                break

        if matched_alias:
            found.append({
                "skill": canonical_name,
                "matched_text": matched_alias,
                "source": "explicit_dictionary_match",
            })

    return found


# ---------------------------------------------------------
# 4. Semantic suggestions using a pretrained model
# These are suggestions, NOT confirmed skills.
# ---------------------------------------------------------

def split_into_sentences(text: str) -> list:
    cleaned = re.sub(r"\s+", " ", text).strip()

    # Simple sentence splitting suitable for a prototype.
    return [
        part.strip()
        for part in re.split(r"(?<=[.!?])\s+|[;\n]", cleaned)
        if len(part.strip()) >= 12
    ]


def semantic_suggestions(
    text: str,
    model: SentenceTransformer,
    explicit_skill_names: set,
    top_k: int = 8,
    threshold: float = 0.48,
) -> list:
    sentences = split_into_sentences(text)

    if not sentences:
        return []

    # Descriptions help the model compare meaning, not just exact words.
    descriptions = {
        skill: f"Professional skill or experience in {skill}."
        for skill in SKILLS
        if skill not in explicit_skill_names
    }

    if not descriptions:
        return []

    skill_names = list(descriptions.keys())
    skill_texts = list(descriptions.values())

    sentence_embeddings = model.encode(
        sentences, convert_to_tensor=True, normalize_embeddings=True
    )
    skill_embeddings = model.encode(
        skill_texts, convert_to_tensor=True, normalize_embeddings=True
    )

    similarities = util.cos_sim(
        sentence_embeddings, skill_embeddings
    )

    best_by_skill = {}

    for skill_index, skill in enumerate(skill_names):
        best_sentence_index = int(
            similarities[:, skill_index].argmax().item()
        )
        score = float(
            similarities[best_sentence_index, skill_index].item()
        )

        if score >= threshold:
            best_by_skill[skill] = {
                "skill": skill,
                "similarity": round(score, 3),
                "evidence_sentence": sentences[best_sentence_index],
                "source": "semantic_candidate_not_verified",
            }

    return sorted(
        best_by_skill.values(),
        key=lambda item: item["similarity"],
        reverse=True,
    )[:top_k]


# ---------------------------------------------------------
# 5. Process one resume
# ---------------------------------------------------------

def process_resume(path: Path, model: SentenceTransformer) -> dict:
    text = read_resume(path)

    if not text.strip():
        return {
            "file": path.name,
            "error": (
                "No text extracted. This may be a scanned PDF; "
                "OCR is not included in this prototype."
            ),
        }

    explicit = extract_explicit_skills(text)
    explicit_names = {item["skill"] for item in explicit}

    suggestions = semantic_suggestions(
        text=text,
        model=model,
        explicit_skill_names=explicit_names,
    )

    # The requested output: an array of canonical skill strings.
    skill_array = [item["skill"] for item in explicit]

    return {
        "file": path.name,
        "text_characters": len(text),
        "skills": skill_array,
        "explicit_evidence": explicit,
        "semantic_candidates": suggestions,
        "warning": (
            "Resume mentions are not proof of ability. "
            "Semantic candidates require human review and assessment."
        ),
    }


# ---------------------------------------------------------
# 6. Accept a file or a folder of resumes
# ---------------------------------------------------------

def main():
    if len(sys.argv) >= 2:
        input_path = Path(sys.argv[1]).expanduser()
    else:
        input_path = Path(
            input("Enter a resume file or folder path: ").strip().strip('"')
        ).expanduser()

    if not input_path.exists():
        print(f"Path not found: {input_path}")
        sys.exit(1)

    if input_path.is_file():
        files = [input_path]
    else:
        supported = {".pdf", ".docx", ".txt"}
        files = sorted(
            p for p in input_path.rglob("*")
            if p.is_file() and p.suffix.lower() in supported
        )

    if not files:
        print("No PDF, DOCX or TXT resumes found.")
        sys.exit(1)

    print("Loading pretrained semantic model...")
    print("The first run downloads model files; later runs reuse the cache.")

    model = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")

    results = []

    for path in files:
        print(f"Processing: {path.name}")

        try:
            results.append(process_resume(path, model))
        except Exception as exc:
            results.append({
                "file": path.name,
                "error": str(exc),
            })

    output = {
        "model": "sentence-transformers/all-MiniLM-L6-v2",
        "resume_count": len(results),
        "results": results,
    }

    output_path = Path("resume_skill_results.json")
    output_path.write_text(
        json.dumps(output, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print("\n========== RESULTS ==========")

    for result in results:
        print(f"\nResume: {result['file']}")

        if "error" in result:
            print(f"Error: {result['error']}")
            continue

        print("Explicit skills:")
        print(json.dumps(result["skills"], indent=2))

        print("Semantic candidates (review these):")
        print(json.dumps(result["semantic_candidates"], indent=2))

    print(f"\nFull results saved to: {output_path.resolve()}")


if __name__ == "__main__":
    main()
